# -*- coding: utf-8 -*-
"""Generate contacts.xlsx — IT helpdesk contact directory (100 rows).
Rerun anytime:  python generate_contacts.py
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

# Department -> (point of contact, role, email, Teams channel, phone)
DEPARTMENTS = {
    "P.Công nghệ Thông tin": ("Nguyễn Văn An", "Trưởng phòng IT", "an.nv@bank.com.vn", "Teams: IT Support", "028-3821-1001"),
    "P.Hạ tầng & Mạng": ("Trần Thị Bình", "Trưởng phòng Hạ tầng", "binh.tt@bank.com.vn", "Teams: Network Team", "028-3821-1002"),
    "P.Bảo mật Thông tin": ("Lê Hoàng Cường", "Trưởng phòng Bảo mật", "cuong.lh@bank.com.vn", "Teams: InfoSec Team", "028-3821-1003"),
    "P.Hệ thống Core Banking": ("Phạm Minh Dũng", "Trưởng phòng Core", "dung.pm@bank.com.vn", "Teams: Core Banking", "028-3821-1004"),
    "P.Thanh toán Quốc tế": ("Vũ Thanh Hà", "Trưởng phòng TTQT", "ha.vt@bank.com.vn", "Teams: SWIFT Team", "028-3821-1005"),
    "P.ATM & Thẻ": ("Đặng Ngọc Em", "Trưởng phòng Thẻ", "em.dn@bank.com.vn", "Teams: Card & ATM", "028-3821-1006"),
    "P.Kế toán": ("Bùi Thị Phượng", "Trưởng phòng Kế toán", "phuong.bt@bank.com.vn", "Teams: Accounting", "028-3821-1007"),
    "P.Nhân sự": ("Hồ Văn Giang", "Trưởng phòng NS", "giang.hv@bank.com.vn", "Teams: HR Team", "028-3821-1008"),
    "P.Dữ liệu & BI": ("Mai Thị Hồng", "Trưởng phòng Dữ liệu", "hong.mt@bank.com.vn", "Teams: Data & BI", "028-3821-1009"),
    "P.Phát triển Phần mềm": ("Phan Quang Khôi", "Trưởng phòng Dev", "khoi.pq@bank.com.vn", "Teams: Dev Team", "028-3821-1010"),
    "P.Vận hành": ("Lương Thanh Long", "Trưởng phòng Vận hành", "long.lt@bank.com.vn", "Teams: Ops Team", "028-3821-1011"),
}

# (nhóm vấn đề, từ khoá nhân viên hay nói, mô tả, phòng ban)
ROWS = [
    # === Máy tính / Hardware (10) ===
    ("Máy tính", "máy tính không bật lỗi khởi động", "Máy trạm/không nguồn, không khởi động được", "P.Công nghệ Thông tin"),
    ("Máy tính", "máy tính chậm đơ treo lag", "Máy chạy chậm, đơ, treo khi dùng", "P.Công nghệ Thông tin"),
    ("Máy tính", "laptop hỏng màn hình đen", "Laptop không lên màn hình", "P.Công nghệ Thông tin"),
    ("Máy tính", "màn hình không hiển thị", "Màn hình desktop không ra hình", "P.Công nghệ Thông tin"),
    ("Máy tính", "bàn phím chuột hỏng", "Bàn phím/chuột không hoạt động", "P.Công nghệ Thông tin"),
    ("Máy tính", "ổ cứng đầy hết dung lượng", "Disk full, không lưu được file", "P.Công nghệ Thông tin"),
    ("Máy tính", "RAM lỗi báo beep", "RAM hỏng, máy báo tiếng beep", "P.Hạ tầng & Mạng"),
    ("Máy tính", "cài đặt phần mềm trên máy", "Yêu cầu cài phần mềm mới", "P.Công nghệ Thông tin"),
    ("Máy tính", "cấp phát máy mới đổi máy", "Xin máy mới / đổi máy tính", "P.Công nghệ Thông tin"),
    ("Máy tính", "bảo trì vệ sinh máy tính", "Bảo trì định kỳ máy trạm", "P.Công nghệ Thông tin"),
    # === Mạng / Network (8) ===
    ("Mạng", "mất mạng không vào internet", "Mất kết nối internet hoàn toàn", "P.Hạ tầng & Mạng"),
    ("Mạng", "wifi chậm không kết nối", "WiFi yếu/chậm/không vào được", "P.Hạ tầng & Mạng"),
    ("Mạng", "vpn không vào lỗi đăng nhập", "VPN không kết nối / sai pass", "P.Bảo mật Thông tin"),
    ("Mạng", "lan cáp mạng đứt", "Cáp mạng/LAN đứt, mất tín hiệu", "P.Hạ tầng & Mạng"),
    ("Mạng", "ip conflict trùng ip", "Trùng địa chỉ IP trên mạng", "P.Hạ tầng & Mạng"),
    ("Mạng", "tốc độ mạng chậm bandwidth", "Băng thông thấp, mạng chậm", "P.Hạ tầng & Mạng"),
    ("Mạng", "firewall chặn truy cập", "Firewall block website/ứng dụng", "P.Bảo mật Thông tin"),
    ("Mạng", "dns lỗi không phân giải", "DNS không phân giải được tên miền", "P.Hạ tầng & Mạng"),
    # === Email / Outlook (6) ===
    ("Email", "không gửi nhận email", "Email không gửi/nhận được", "P.Công nghệ Thông tin"),
    ("Email", "email đầy hết quota", "Hộp thư đầy, vượt dung lượng", "P.Công nghệ Thông tin"),
    ("Email", "outlook lỗi không mở", "Outlook crash/không mở được", "P.Công nghệ Thông tin"),
    ("Email", "quên mật khẩu email", "Reset mật khẩu email", "P.Bảo mật Thông tin"),
    ("Email", "email spam phishing lừa đảo", "Nhận email nghi ngờ lừa đảo", "P.Bảo mật Thông tin"),
    ("Email", "cấu hình email điện thoại", "Cài email trên mobile", "P.Công nghệ Thông tin"),
    # === SWIFT / Thanh toán quốc tế (10) ===
    ("SWIFT", "swift message bị NAK", "Tin nhắn SWIFT bị từ chối (NAK)", "P.Thanh toán Quốc tế"),
    ("SWIFT", "swift ACK chậm", "ACK xác nhận SWIFT chậm", "P.Thanh toán Quốc tế"),
    ("SWIFT", "swift kết nối mất SAG lỗi", "SWIFT Alliance Gateway mất kết nối", "P.Thanh toán Quốc tế"),
    ("SWIFT", "chuyển tiền quốc tế lỗi", "Lệnh chuyển tiền QT bị lỗi", "P.Thanh toán Quốc tế"),
    ("SWIFT", "swift key BIC code sai", "BIC/key sai, tin nhắn bị reject", "P.Thanh toán Quốc tế"),
    ("SWIFT", "swift message duplicate trùng", "Tin nhắn SWIFT bị trùng/gửi 2 lần", "P.Thanh toán Quốc tế"),
    ("SWIFT", "reversal hủy lệnh swift", "Hủy/reversal lệnh SWIFT", "P.Thanh toán Quốc tế"),
    ("SWIFT", "swift queue stuck kẹt hàng", "Hàng đợi SWIFT bị kẹt", "P.Thanh toán Quốc tế"),
    ("SWIFT", "đối soát swift reconciliation", "Đối soát giao dịch SWIFT", "P.Thanh toán Quốc tế"),
    ("SWIFT", "cấu hình swift alliance access", "Cấu hình SWIFT Alliance Access", "P.Thanh toán Quốc tế"),
    # === Core Banking (8) ===
    ("Core Banking", "core banking chậm treo", "Hệ thống core chậm/treo", "P.Hệ thống Core Banking"),
    ("Core Banking", "core không đăng nhập được", "Không login được core banking", "P.Hệ thống Core Banking"),
    ("Core Banking", "giao dịch core lỗi fail", "Giao dịch trên core bị fail", "P.Hệ thống Core Banking"),
    ("Core Banking", "cập nhật số dư chậm", "Số dư cập nhật chậm/bất đồng bộ", "P.Hệ thống Core Banking"),
    ("Core Banking", "batch job core fail", "Batch job core banking thất bại", "P.Vận hành"),
    ("Core Banking", "restore backup core", "Khôi phục/sao lưu core banking", "P.Hệ thống Core Banking"),
    ("Core Banking", "báo cáo core lỗi", "Báo cáo từ core bị lỗi/sai số liệu", "P.Hệ thống Core Banking"),
    ("Core Banking", "cấu hình sản phẩm core", "Cấu hình sản phẩm/lãi suất trên core", "P.Hệ thống Core Banking"),
    # === ATM / Thẻ (8) ===
    ("ATM & Thẻ", "atm nuốt thẻ kẹt thẻ", "ATM nuốt/kẹt thẻ khách", "P.ATM & Thẻ"),
    ("ATM & Thẻ", "atm không trả tiền lỗi rút", "ATM không xuất tiền", "P.ATM & Thẻ"),
    ("ATM & Thẻ", "thẻ bị khóa block", "Thẻ bị khóa/block", "P.ATM & Thẻ"),
    ("ATM & Thẻ", "thẻ hết hạn gia hạn", "Thẻ hết hạn, cần gia hạn", "P.ATM & Thẻ"),
    ("ATM & Thẻ", "pos lỗi không quẹt", "Máy POS không quẹt thẻ được", "P.ATM & Thẻ"),
    ("ATM & Thẻ", "quên pin khóa pin", "Quên/khóa mã PIN thẻ", "P.ATM & Thẻ"),
    ("ATM & Thẻ", "giao dịch thẻ trừ sai", "Thẻ bị trừ tiền sai số", "P.ATM & Thẻ"),
    ("ATM & Thẻ", "kích hoạt thẻ mới", "Kích hoạt thẻ mới phát hành", "P.ATM & Thẻ"),
    # === Bảo mật / Security (8) ===
    ("Bảo mật", "tài khoản bị khóa lock", "Tài khoản AD/email bị khóa", "P.Bảo mật Thông tin"),
    ("Bảo mật", "quên mật khẩu reset password", "Yêu cầu reset mật khẩu", "P.Bảo mật Thông tin"),
    ("Bảo mật", "2fa otp không nhận", "Không nhận OTP/2FA", "P.Bảo mật Thông tin"),
    ("Bảo mật", "nghi ngờ hack xâm nhập", "Nghi ngờ tài khoản bị xâm nhập", "P.Bảo mật Thông tin"),
    ("Bảo mật", "virus malware máy", "Máy nhiễm virus/malware", "P.Bảo mật Thông tin"),
    ("Bảo mật", "email lừa đảo phishing", "Báo cáo email phishing", "P.Bảo mật Thông tin"),
    ("Bảo mật", "usb thiết bị bị chặn", "USB/thiết bị ngoài bị policy chặn", "P.Bảo mật Thông tin"),
    ("Bảo mật", "token RSA hard token lỗi", "Token RSA/hard token không hoạt động", "P.Bảo mật Thông tin"),
    # === Phần mềm / Software (10) ===
    ("Phần mềm", "office word excel lỗi", "Word/Excel/PowerPoint lỗi", "P.Công nghệ Thông tin"),
    ("Phần mềm", "phần mềm ERP lỗi", "Hệ thống ERP bị lỗi", "P.Phát triển Phần mềm"),
    ("Phần mềm", "phần mềm CRM lỗi", "Hệ thống CRM bị lỗi", "P.Phát triển Phần mềm"),
    ("Phần mềm", "phần mềm kế toán SAP lỗi", "SAP/phần mềm kế toán lỗi", "P.Kế toán"),
    ("Phần mềm", "cài đặt phần mềm mới", "Yêu cầu cài phần mềm", "P.Công nghệ Thông tin"),
    ("Phần mềm", "license phần mềm hết hạn", "License hết hạn cần gia hạn", "P.Công nghệ Thông tin"),
    ("Phần mềm", "phần mềm không mở crash", "Phần mềm crash/không mở được", "P.Công nghệ Thông tin"),
    ("Phần mềm", "cập nhật phần mềm", "Update phiên bản phần mềm", "P.Công nghệ Thông tin"),
    ("Phần mềm", "tích hợp API hệ thống", "Tích hợp/API giữa các hệ thống", "P.Phát triển Phần mềm"),
    ("Phần mềm", "lỗi version compatibility", "Xung đột phiên bản phần mềm", "P.Phát triển Phần mềm"),
    # === Hệ thống / Server / Database (8) ===
    ("Server/DB", "server down sập", "Server sập/không truy cập", "P.Hạ tầng & Mạng"),
    ("Server/DB", "database lỗi không truy cập", "DB không kết nối được", "P.Dữ liệu & BI"),
    ("Server/DB", "database chậm query timeout", "Query chậm/timeout", "P.Dữ liệu & BI"),
    ("Server/DB", "backup fail không backup", "Backup thất bại", "P.Vận hành"),
    ("Server/DB", "restore dữ liệu", "Khôi phục dữ liệu", "P.Vận hành"),
    ("Server/DB", "disk full server", "Disk server đầy", "P.Hạ tầng & Mạng"),
    ("Server/DB", "cpu ram server cao", "Server quá tải CPU/RAM", "P.Hạ tầng & Mạng"),
    ("Server/DB", "log server đầy", "Log file đầy ổ cứng", "P.Vận hành"),
    # === Website / Portal / Mobile (6) ===
    ("Website/App", "website không vào được", "Website ngân hàng không truy cập", "P.Phát triển Phần mềm"),
    ("Website/App", "internet banking lỗi", "IBanking lỗi/không vào được", "P.Phát triển Phần mềm"),
    ("Website/App", "mobile app lỗi crash", "App mobile crash/lỗi", "P.Phát triển Phần mềm"),
    ("Website/App", "portal nhân sự không vào", "Portal HR không truy cập", "P.Nhân sự"),
    ("Website/App", "lỗi đăng nhập portal", "Không login được portal", "P.Bảo mật Thông tin"),
    ("Website/App", "cập nhật nội dung website", "Đăng/sửa nội dung website", "P.Phát triển Phần mềm"),
    # === HR / Nhân sự (5) ===
    ("Nhân sự", "phần mềm chấm công lỗi", "Chấm công lỗi/không vào được", "P.Nhân sự"),
    ("Nhân sự", "email nhân sự hợp đồng", "Hỏi về hợp đồng/chính sách NS", "P.Nhân sự"),
    ("Nhân sự", "cấp tài khoản nhân viên mới", "Tạo tài khoản cho NV mới", "P.Công nghệ Thông tin"),
    ("Nhân sự", "khóa tài khoản nghỉ việc", "Khóa tài khoản NV nghỉ việc", "P.Bảo mật Thông tin"),
    ("Nhân sự", "đổi thông tin cá nhân", "Cập nhật thông tin cá nhân", "P.Nhân sự"),
    # === Kế toán / Finance (4) ===
    ("Kế toán", "phần mềm kế toán lỗi", "Phần mềm KT lỗi/không vào", "P.Kế toán"),
    ("Kế toán", "báo cáo tài chính sai", "Báo cáo TC/số liệu sai", "P.Kế toán"),
    ("Kế toán", "duyệt chi thanh toán nội bộ", "Duyệt chi/thanh toán nội bộ", "P.Kế toán"),
    ("Kế toán", "hóa đơn điện tử lỗi", "HĐĐT lỗi/không xuất được", "P.Kế toán"),
    # === Cloud / VNG Cloud (4) ===
    ("Cloud", "vng cloud server không truy cập", "Máy chủ VNG Cloud không vào", "P.Hạ tầng & Mạng"),
    ("Cloud", "tạo xóa máy chủ cloud", "Yêu cầu tạo/xóa server cloud", "P.Hạ tầng & Mạng"),
    ("Cloud", "cloud billing chi phí", "Hỏi chi phí/billing cloud", "P.Kế toán"),
    ("Cloud", "object storage s3 lỗi", "Object storage/S3 lỗi", "P.Hạ tầng & Mạng"),
    # === In ấn / Scanner / VoIP (5) ===
    ("In/Scan/VoIP", "máy in lỗi không in", "Máy in không in được", "P.Công nghệ Thông tin"),
    ("In/Scan/VoIP", "máy in kẹt giấy", "Máy in kẹt giấy", "P.Công nghệ Thông tin"),
    ("In/Scan/VoIP", "scanner máy scan lỗi", "Máy scan không hoạt động", "P.Công nghệ Thông tin"),
    ("In/Scan/VoIP", "điện thoại bàn voip lỗi", "ĐT bàn/VoIP không gọi được", "P.Hạ tầng & Mạng"),
    ("In/Scan/VoIP", "fax lỗi", "Máy fax không gửi/nhận", "P.Vận hành"),
]

def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dau_moi"
    headers = ["STT", "Nhom_van_de", "Tu_khoa", "Mo_ta", "Phong_ban",
               "Nguoi_dau_moi", "Chuc_vu", "Email", "Teams", "Dien_thoai"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, (nhom, tukhoa, mota, phong) in enumerate(ROWS, start=1):
        nguoi, chucvu, email, teams, dt = DEPARTMENTS[phong]
        ws.append([i, nhom, tukhoa, mota, phong, nguoi, chucvu, email, teams, dt])
    widths = [5, 16, 34, 40, 26, 18, 20, 26, 20, 16]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    out = "contacts.xlsx"
    wb.save(out)
    print(f"Created {out} with {len(ROWS)} rows.")

if __name__ == "__main__":
    main()
